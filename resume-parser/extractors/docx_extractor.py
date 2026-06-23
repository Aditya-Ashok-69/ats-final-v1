from docx import Document
from docx.opc.exceptions import PackageNotFoundError


class DocxExtractionError(Exception):
    """Raised when a DOCX cannot be opened or contains no extractable text."""
    pass


def extract_docx_text(file_path: str) -> str:
    """
    Extract text from a DOCX file, including text inside tables
    (resumes frequently use tables for layout, and the original
    version silently dropped that content).

    Raises:
        DocxExtractionError: if the file is corrupt or has no text.
    """
    try:
        doc = Document(file_path)
    except PackageNotFoundError:
        raise DocxExtractionError("File is not a valid .docx (corrupt or wrong format)")
    except Exception as e:
        raise DocxExtractionError(f"Could not open DOCX: {e}")

    parts = [p.text for p in doc.paragraphs if p.text.strip()]

    # Pull text out of tables too -- common in resume templates.
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)

    text = "\n".join(parts).strip()

    if not text:
        raise DocxExtractionError("No extractable text found in DOCX")

    return text
